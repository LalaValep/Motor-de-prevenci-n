#!/usr/bin/env python3
import json
import os
import argparse
import datetime
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter

DEFAULT_LOG    = '/tmp/security_events.log'
DEFAULT_RULES  = os.path.join(os.path.dirname(__file__), 'security_rules.json')
DEFAULT_OUTPUT = os.path.expanduser(
    '~/Downloads/Proyecto de grado/otras/informe_amenazas.xlsx'
)

# Colores
COLOR_HEADER    = '1F3864'  # Azul oscuro
COLOR_SUBHEADER = '2E75B6'  # Azul medio
COLOR_CRITICO   = 'C00000'  # Rojo
COLOR_ALTO      = 'FF0000'  # Rojo claro
COLOR_SOSPECHA  = 'FF9900'  # Naranja
COLOR_OK        = '70AD47'  # Verde
COLOR_ROW_ALT   = 'D9E1F2'  # Azul claro alternado
COLOR_WHITE     = 'FFFFFF'
COLOR_GOLD      = 'FFD700'

def side():
    return Side(style='thin', color='BFBFBF')

def border():
    s = side()
    return Border(left=s, right=s, top=s, bottom=s)

def header_font(size=11):
    return Font(name='Arial', bold=True, color=COLOR_WHITE, size=size)

def normal_font(size=10):
    return Font(name='Arial', size=size)

def bold_font(size=10):
    return Font(name='Arial', bold=True, size=size)

def fill(color):
    return PatternFill('solid', start_color=color, fgColor=color)

def center():
    return Alignment(horizontal='center', vertical='center', wrap_text=True)

def left():
    return Alignment(horizontal='left', vertical='center', wrap_text=True)

def set_col_width(ws, col, width):
    ws.column_dimensions[get_column_letter(col)].width = width

def style_header_row(ws, row, cols, bg=COLOR_HEADER):
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font      = header_font()
        cell.fill      = fill(bg)
        cell.alignment = center()
        cell.border    = border()

def style_data_row(ws, row, cols, alt=False):
    bg = COLOR_ROW_ALT if alt else COLOR_WHITE
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font      = normal_font()
        cell.fill      = fill(bg)
        cell.alignment = left()
        cell.border    = border()

def load_rules(path):
    try:
        with open(path, 'r', encoding='utf-8') as f:
            data = json.load(f)
        index = {}
        for r in data.get('reglas', []):
            index[r['ataque_cibernetico'].lower()] = r
        return index
    except Exception as e:
        print(f'[ERROR] Reglas: {e}')
        return {}

def load_events(path):
    events = []
    try:
        with open(path, 'r', encoding='utf-8') as f:
            for line in f:
                line = line.strip()
                if line:
                    try:
                        events.append(json.loads(line))
                    except:
                        pass
        return events
    except Exception as e:
        print(f'[ERROR] Log: {e}')
        return []

def get_rule(evento, rules):
    tipo = evento.get('tipo_ataque', '').lower()
    mapeo = {
        'arp_spoofing':         'suplantación (spoofing)',
        'mitm':                 'ataque de intermediario o reproducción (mitm)',
        'mitm_puerto_sensible': 'ataque de intermediario o reproducción (mitm)',
        'malware_distribution': 'distribución de malware',
    }
    clave = mapeo.get(tipo)
    return rules.get(clave) if clave else None

def analyze(events):
    r = {
        'total': len(events),
        'por_tipo': defaultdict(int),
        'por_nivel': defaultdict(int),
        'por_accion': defaultdict(int),
        'por_atacante': defaultdict(list),
        'confirmados': [],
        'altos': [],
        'leves': [],
        'falsos': [],
        'victimas': set(),
        'atacantes': set(),
    }
    for e in events:
        tipo   = e.get('tipo_ataque', 'DESCONOCIDO')
        nivel  = e.get('nivel_sospecha', 'DESCONOCIDO')
        accion = e.get('accion_tomada', 'DESCONOCIDO')
        mac    = e.get('mac_atacante', 'N/A')
        r['por_tipo'][tipo]    += 1
        r['por_nivel'][nivel]  += 1
        r['por_accion'][accion]+= 1
        r['por_atacante'][mac].append(e)
        r['atacantes'].add(mac)
        for v in e.get('victimas', []):
            r['victimas'].add(v)
        if tipo == 'FALSO_POSITIVO':
            r['falsos'].append(e)
        elif nivel == 'ATAQUE_CONFIRMADO':
            r['confirmados'].append(e)
        elif nivel == 'SOSPECHA_ALTA':
            r['altos'].append(e)
        elif nivel == 'SOSPECHA_LEVE':
            r['leves'].append(e)
    return r

def sheet_resumen(wb, events, res, rules):
    ws = wb.active
    ws.title = 'Resumen Ejecutivo'
    ws.sheet_view.showGridLines = False

    # Título principal
    ws.merge_cells('A1:F1')
    c = ws['A1']
    c.value     = 'INFORME DE AMENAZAS — MOTOR DE SEGURIDAD SDN'
    c.font      = Font(name='Arial', bold=True, color=COLOR_WHITE, size=14)
    c.fill      = fill(COLOR_HEADER)
    c.alignment = center()
    ws.row_dimensions[1].height = 30

    ws.merge_cells('A2:F2')
    c = ws['A2']
    c.value     = 'Proyecto: Seguridad en Redes Wi-Fi Públicas Simuladas'
    c.font      = Font(name='Arial', color=COLOR_WHITE, size=11)
    c.fill      = fill(COLOR_SUBHEADER)
    c.alignment = center()

    ws.merge_cells('A3:F3')
    c = ws['A3']
    c.value     = f'Generado: {datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")}'
    c.font      = normal_font()
    c.alignment = center()
    c.fill      = fill(COLOR_ROW_ALT)

    # Métricas generales
    row = 5
    ws.merge_cells(f'A{row}:F{row}')
    c = ws.cell(row=row, column=1, value='MÉTRICAS GENERALES')
    c.font = header_font()
    c.fill = fill(COLOR_SUBHEADER)
    c.alignment = center()
    style_header_row(ws, row, 6, COLOR_SUBHEADER)

    metricas = [
        ('Total de eventos analizados', res['total']),
        ('Atacantes únicos detectados', len(res['atacantes'])),
        ('Víctimas únicas afectadas',   len(res['victimas'])),
        ('Ataques confirmados',          len(res['confirmados'])),
        ('Sospechas altas',              len(res['altos'])),
        ('Sospechas leves',              len(res['leves'])),
        ('Falsos positivos',             len(res['falsos'])),
    ]

    for i, (label, val) in enumerate(metricas):
        r = row + 1 + i
        ws.cell(row=r, column=1, value=label).font = bold_font()
        ws.cell(row=r, column=1).fill = fill(COLOR_ROW_ALT if i % 2 == 0 else COLOR_WHITE)
        ws.cell(row=r, column=1).border = border()
        ws.merge_cells(f'A{r}:D{r}')

        cell_val = ws.cell(row=r, column=5, value=val)
        cell_val.font      = bold_font(12)
        cell_val.alignment = center()
        cell_val.border    = border()

        # Color según tipo
        if 'confirmados' in label.lower():
            cell_val.fill = fill(COLOR_CRITICO)
            cell_val.font = Font(name='Arial', bold=True, color=COLOR_WHITE, size=12)
        elif 'sospechas' in label.lower():
            cell_val.fill = fill(COLOR_SOSPECHA)
        elif 'falsos' in label.lower():
            cell_val.fill = fill(COLOR_OK)
            cell_val.font = Font(name='Arial', bold=True, color=COLOR_WHITE, size=12)
        else:
            cell_val.fill = fill(COLOR_ROW_ALT)
        ws.merge_cells(f'E{r}:F{r}')

    # Eventos por tipo
    row = row + len(metricas) + 2
    ws.merge_cells(f'A{row}:F{row}')
    c = ws.cell(row=row, column=1, value='EVENTOS POR TIPO DE ATAQUE')
    c.font = header_font()
    c.fill = fill(COLOR_SUBHEADER)
    c.alignment = center()

    headers = ['Tipo de Ataque', 'Cantidad', 'Porcentaje']
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=row+1, column=i, value=h)
        c.font = header_font()
        c.fill = fill(COLOR_HEADER)
        c.alignment = center()
        c.border = border()

    total = res['total'] or 1
    for i, (tipo, count) in enumerate(sorted(res['por_tipo'].items(),
                                              key=lambda x: x[1], reverse=True)):
        r = row + 2 + i
        pct = f'{count/total*100:.1f}%'
        style_data_row(ws, r, 3, i % 2 == 0)
        ws.cell(row=r, column=1, value=tipo)
        ws.cell(row=r, column=2, value=count).alignment = center()
        ws.cell(row=r, column=3, value=pct).alignment  = center()

    # Anchos
    set_col_width(ws, 1, 40)
    set_col_width(ws, 2, 15)
    set_col_width(ws, 3, 15)
    set_col_width(ws, 4, 20)
    set_col_width(ws, 5, 15)
    set_col_width(ws, 6, 15)


def sheet_eventos(wb, events, res, rules, titulo, lista, color_nivel):
    ws = wb.create_sheet(titulo)
    ws.sheet_view.showGridLines = False

    ws.merge_cells('A1:H1')
    c = ws['A1']
    c.value     = titulo.upper()
    c.font      = header_font(13)
    c.fill      = fill(COLOR_HEADER)
    c.alignment = center()
    ws.row_dimensions[1].height = 28

    headers = ['#', 'Timestamp', 'Tipo Ataque', 'Nivel Riesgo',
               'Nivel Sospecha', 'MAC Atacante', 'IP Atacante', 'Víctimas', 'Acción', 'Detalle']
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=2, column=i, value=h)
        c.font      = header_font()
        c.fill      = fill(color_nivel)
        c.alignment = center()
        c.border    = border()

    if not lista:
        ws.merge_cells('A3:J3')
        c = ws.cell(row=3, column=1, value='No se registraron eventos en esta categoría.')
        c.font      = Font(name='Arial', italic=True, size=10)
        c.alignment = center()
    else:
        for i, evento in enumerate(lista):
            r    = i + 3
            regla = get_rule(evento, rules)
            style_data_row(ws, r, 10, i % 2 == 0)
            ws.cell(row=r, column=1,  value=i+1).alignment = center()
            ws.cell(row=r, column=2,  value=evento.get('timestamp', ''))
            ws.cell(row=r, column=3,  value=evento.get('tipo_ataque', ''))
            ws.cell(row=r, column=4,  value=regla['nivel_de_riesgo'] if regla else 'N/A')
            ws.cell(row=r, column=5,  value=evento.get('nivel_sospecha', ''))
            ws.cell(row=r, column=6,  value=evento.get('mac_atacante', ''))
            ws.cell(row=r, column=7,  value=evento.get('ip_atacante', ''))
            ws.cell(row=r, column=8,  value=', '.join(evento.get('victimas', [])))
            ws.cell(row=r, column=9,  value=evento.get('accion_tomada', ''))
            ws.cell(row=r, column=10, value=evento.get('detalle', ''))

    widths = [5, 20, 25, 15, 20, 20, 15, 20, 20, 50]
    for i, w in enumerate(widths, 1):
        set_col_width(ws, i, w)


def sheet_contexto(wb, events, rules):
    ws = wb.create_sheet('Contexto de Ataques')
    ws.sheet_view.showGridLines = False

    ws.merge_cells('A1:E1')
    c = ws['A1']
    c.value     = 'CONTEXTO DE ATAQUES — SECURITY_RULES.JSON'
    c.font      = header_font(13)
    c.fill      = fill(COLOR_HEADER)
    c.alignment = center()
    ws.row_dimensions[1].height = 28

    tipos_ocurridos = set(e.get('tipo_ataque') for e in events) - {'FALSO_POSITIVO'}
    reglas_mostradas = set()
    row = 3

    for tipo in tipos_ocurridos:
        ejemplo = next((e for e in events if e.get('tipo_ataque') == tipo), {})
        regla   = get_rule(ejemplo, rules)
        if not regla or regla['ataque_cibernetico'] in reglas_mostradas:
            continue
        reglas_mostradas.add(regla['ataque_cibernetico'])

        # Encabezado del ataque
        ws.merge_cells(f'A{row}:E{row}')
        c = ws.cell(row=row, column=1, value=f'Ataque: {regla["ataque_cibernetico"]}')
        c.font = header_font()
        c.fill = fill(COLOR_SUBHEADER)
        c.alignment = left()
        row += 1

        ws.merge_cells(f'A{row}:E{row}')
        c = ws.cell(row=row, column=1, value=f'Nivel de Riesgo: {regla["nivel_de_riesgo"]}')
        c.font = bold_font()
        c.fill = fill(COLOR_CRITICO if 'CRÍTICO' in regla['nivel_de_riesgo'] else COLOR_SOSPECHA)
        c.font = Font(name='Arial', bold=True, color=COLOR_WHITE)
        c.alignment = left()
        row += 1

        # Factores
        ws.merge_cells(f'A{row}:E{row}')
        ws.cell(row=row, column=1, value='Factores de Vulnerabilidad Asociados').font = bold_font()
        ws.cell(row=row, column=1).fill = fill(COLOR_ROW_ALT)
        row += 1

        headers = ['Factor', 'Descripción']
        for i, h in enumerate(headers, 1):
            c = ws.cell(row=row, column=i, value=h)
            c.font      = header_font()
            c.fill      = fill(COLOR_HEADER)
            c.alignment = center()
            c.border    = border()
        row += 1

        for j, f_item in enumerate(regla.get('factores_asociados_a_la_vulnerabilidad', [])):
            style_data_row(ws, row, 2, j % 2 == 0)
            ws.cell(row=row, column=1, value=f_item['factor'])
            ws.cell(row=row, column=2, value=f_item['descripcion'])
            row += 1

        # Soluciones
        row += 1
        ws.merge_cells(f'A{row}:E{row}')
        ws.cell(row=row, column=1, value='Soluciones Aplicadas por el Controlador').font = bold_font()
        ws.cell(row=row, column=1).fill = fill(COLOR_ROW_ALT)
        row += 1

        headers2 = ['Solución', 'Acción del Controlador']
        for i, h in enumerate(headers2, 1):
            c = ws.cell(row=row, column=i, value=h)
            c.font      = header_font()
            c.fill      = fill(COLOR_HEADER)
            c.alignment = center()
            c.border    = border()
        row += 1

        for j, sol in enumerate(regla.get('soluciones_comunmente_usadas', [])):
            style_data_row(ws, row, 2, j % 2 == 0)
            ws.cell(row=row, column=1, value=sol['solucion'])
            ws.cell(row=row, column=2, value=sol['accion_controlador'])
            row += 1

        row += 2

    set_col_width(ws, 1, 45)
    set_col_width(ws, 2, 70)


def sheet_hosts(wb, res, rules):
    ws = wb.create_sheet('Hosts Anómalos')
    ws.sheet_view.showGridLines = False

    ws.merge_cells('A1:F1')
    c = ws['A1']
    c.value     = 'HOSTS CON MAYOR ACTIVIDAD ANÓMALA'
    c.font      = header_font(13)
    c.fill      = fill(COLOR_HEADER)
    c.alignment = center()
    ws.row_dimensions[1].height = 28

    headers = ['MAC Atacante', 'Total Eventos', 'Tipos Detectados',
               'Última Acción', 'Último Nivel', 'Estado']
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=2, column=i, value=h)
        c.font      = header_font()
        c.fill      = fill(COLOR_SUBHEADER)
        c.alignment = center()
        c.border    = border()

    ordenados = sorted(res['por_atacante'].items(),
                       key=lambda x: len(x[1]), reverse=True)

    for i, (mac, evts) in enumerate(ordenados):
        if mac == 'N/A':
            continue
        r      = i + 3
        tipos  = ', '.join(set(e['tipo_ataque'] for e in evts))
        ultima = evts[-1]
        nivel  = ultima.get('nivel_sospecha', 'N/A')
        estado = 'BLOQUEADO' if nivel == 'ATAQUE_CONFIRMADO' else 'EN MONITOREO'

        style_data_row(ws, r, 6, i % 2 == 0)
        ws.cell(row=r, column=1, value=mac)
        ws.cell(row=r, column=2, value=len(evts)).alignment = center()
        ws.cell(row=r, column=3, value=tipos)
        ws.cell(row=r, column=4, value=ultima.get('accion_tomada', 'N/A'))
        ws.cell(row=r, column=5, value=nivel)

        c_estado = ws.cell(row=r, column=6, value=estado)
        c_estado.alignment = center()
        c_estado.font = Font(name='Arial', bold=True, color=COLOR_WHITE)
        c_estado.fill = fill(COLOR_CRITICO if estado == 'BLOQUEADO' else COLOR_SOSPECHA)

    widths = [22, 15, 40, 22, 22, 18]
    for i, w in enumerate(widths, 1):
        set_col_width(ws, i, w)

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--log',   default=DEFAULT_LOG)
    parser.add_argument('--rules', default=DEFAULT_RULES)
    parser.add_argument('--out',   default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    print('\n' + '='*55)
    print('  Generador de Informe de Amenazas — Excel')
    print('='*55)

    rules  = load_rules(args.rules)
    events = load_events(args.log)

    if not events:
        print('\nNo hay eventos. Ejecuta la simulación primero.')
        return

    res = analyze(events)
    wb  = Workbook()

    sheet_resumen(wb, events, res, rules)
    sheet_eventos(wb, events, res, rules,
                  'Ataques Confirmados', res['confirmados'], COLOR_CRITICO)
    sheet_eventos(wb, events, res, rules,
                  'Sospechas Altas', res['altos'], COLOR_SOSPECHA)
    sheet_eventos(wb, events, res, rules,
                  'Sospechas Leves', res['leves'], '4472C4')
    sheet_eventos(wb, events, res, rules,
                  'Falsos Positivos', res['falsos'], COLOR_OK)
    sheet_hosts(wb, res, rules)
    sheet_contexto(wb, events, rules)

    # Crear carpeta si no existe
    os.makedirs(os.path.dirname(args.out), exist_ok=True)
    wb.save(args.out)
    print(f'\n[OK] Informe guardado en:\n     {args.out}\n')


if __name__ == '__main__':
    main()